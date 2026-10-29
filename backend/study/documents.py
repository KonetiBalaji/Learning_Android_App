"""Bounded document extraction and deterministic structure detection."""
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_BYTES = 50 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_PAGES = 2000


class DocumentError(ValueError):
    pass


def extract(filename: str, content: bytes) -> list[str]:
    if not content or len(content) > MAX_BYTES:
        raise DocumentError("Choose a nonempty file up to 50 MB.")
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")
            if "\x00" in text:
                raise DocumentError("TXT must contain UTF-8 text, not binary data.")
            pages = [text]
        elif extension == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise DocumentError("This file is not a valid PDF.")
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise DocumentError("Password-protected PDFs are not supported. Upload an unlocked copy.")
            if len(reader.pages) > MAX_PAGES:
                raise DocumentError("PDF exceeds the 2,000-page processing limit.")
            pages = []
            total = 0
            for page in reader.pages:
                text = page.extract_text() or ""
                total += len(text)
                if total > MAX_TEXT:
                    raise DocumentError("Extracted text exceeds the 2 million character limit.")
                pages.append(text)
        elif extension == ".docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise DocumentError("DOCX expands beyond the processing limit.")
                if "word/document.xml" not in archive.namelist():
                    raise DocumentError("This file is not a valid DOCX.")
            doc = Document(io.BytesIO(content))
            paragraphs = []
            for paragraph in doc.paragraphs:
                style = paragraph.style.name if paragraph.style else ""
                if style.startswith("Heading "):
                    prefix = "# " if style == "Heading 1" else "## "
                    paragraphs.append(prefix + paragraph.text)
                else:
                    paragraphs.append(paragraph.text)
            paragraphs.extend(" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
            pages = ["\n".join(paragraphs)]
        else:
            raise DocumentError("Supported formats: PDF, DOCX, and UTF-8 TXT.")
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError("Could not read this document. It may be corrupted or use an unsupported encoding.") from error
    if sum(map(len, pages)) > MAX_TEXT:
        raise DocumentError("Extracted text exceeds the 2 million character limit.")
    if not any(p.strip() for p in pages):
        raise DocumentError("No readable text found. Scanned documents need OCR before upload.")
    return pages


def structure(filename: str, pages: list[str]) -> dict:
    sections = []
    passages = []
    current = "s0"
    chapter = None
    sections.append({"id": current, "title": "Opening / unsectioned material", "parent_id": None})
    heading = re.compile(r"^(?:chapter\s+\w+|section\s+\w+|\d+(?:\.\d+)*[.)]?\s+\S).{0,100}$", re.I)
    for page_number, page in enumerate(pages, 1):
        buffer = []

        def flush():
            if not buffer:
                return
            paragraph = " ".join(buffer)
            buffer.clear()
            for sentence in re.split(r"(?<=[.!?])\s+", paragraph):
                for start in range(0, len(sentence), 2000):
                    passages.append({"id": f"p{len(passages)}", "page": page_number, "section_id": current, "text": sentence[start:start + 2000]})
                    if len(passages) > 20000:
                        raise DocumentError("Document exceeds the 20,000-passage processing limit.")

        for line in page.splitlines() + [""]:
            line = line.strip()
            if not line:
                flush()
                continue
            is_heading = len(line) < 120 and (heading.match(line) or line.startswith("# ") or line.startswith("## ") or (line.isupper() and len(line.split()) < 12))
            if is_heading:
                flush()
                top_level = bool(re.match(r"chapter\s", line, re.I)) or line.startswith("# ") or (line.isupper() and not re.match(r"section\s", line, re.I))
                current = f"s{len(sections)}"
                sections.append({"id": current, "title": line.lstrip("#").strip(), "parent_id": None if top_level else chapter})
                if top_level:
                    chapter = current
                if len(sections) > 1000:
                    raise DocumentError("Document exceeds the 1,000-section processing limit.")
                continue
            buffer.append(line)
            if re.search(r"[.!?]$", line):
                flush()
    populated = {p["section_id"] for p in passages}
    populated.update(s["parent_id"] for s in sections if s["id"] in populated and s["parent_id"])
    sections = [s for s in sections if s["id"] in populated]
    if not passages:
        raise DocumentError("The document contains headings but no study passages.")
    return {"title": Path(filename).stem, "page_count": len(pages), "has_real_pages": filename.lower().endswith(".pdf"),
            "structure_confidence": "detected" if len(sections) > 1 else "uncertain", "sections": sections, "passages": passages}


def retrieve(document: dict, selection: dict) -> list[dict]:
    passages = document["passages"]
    ids = selection.get("section_ids", [])
    if ids:
        known = {s["id"] for s in document["sections"]}
        if not set(ids) <= known:
            raise DocumentError("A selected section no longer exists.")
        expanded = set(ids)
        expanded.update(s["id"] for s in document["sections"] if s.get("parent_id") in expanded)
        passages = [p for p in passages if p["section_id"] in expanded]
    first, last = selection.get("page_start"), selection.get("page_end")
    if (first is None) != (last is None):
        raise DocumentError("Enter both the first and last page.")
    if first is not None:
        if not document["has_real_pages"] or not 1 <= first <= last <= document["page_count"]:
            raise DocumentError("Choose a valid PDF page range.")
        passages = [p for p in passages if first <= p["page"] <= last]
    topic = selection.get("topic", "").strip()
    if topic:
        stopwords = {"the", "a", "an", "and", "or", "of", "to", "in", "on", "is", "are", "me", "quiz", "what", "about", "only"}
        words = set(re.findall(r"\w+", topic.lower())) - stopwords
        if not words:
            raise DocumentError("Enter a specific topic keyword, such as genetics or energy.")
        matching_sections = {s["id"] for s in document["sections"] if words & set(re.findall(r"\w+", s["title"].lower()))}
        matching_sections.update(s["id"] for s in document["sections"] if s.get("parent_id") in matching_sections)
        passages = [p for p in passages if p["section_id"] in matching_sections or words & set(re.findall(r"\w+", p["text"].lower()))]
    if not passages:
        raise DocumentError("No passages match your selection. Broaden the sections, pages, or topic.")
    return passages
