import io
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from study.api import create_app
from study.documents import DocumentError, extract, retrieve, structure
from study.quizzes import GenerationError, generate_local, validate
from study.storage import Store

MATERIAL = """Chapter 1 Cells
Mitochondria produce usable energy through cellular respiration inside living cells.
Ribosomes assemble proteins using instructions carried by messenger molecules.
Chloroplasts capture sunlight and transform light into chemical energy.
Membranes control the movement of substances into and outside cells.
Chromosomes contain genetic information passed between successive generations.
Vacuoles store water and dissolved substances within plant cells.
Chapter 2 Ecology
Ecosystems include organisms interacting with their surrounding physical environment.
Populations consist of individuals belonging to the same species.
Communities contain several populations living together in one habitat.
Decomposers break down dead material and recycle essential nutrients.
Producers manufacture their own food using sunlight or chemical reactions.
Consumers obtain energy by eating producers or other living organisms.
"""


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path)) as client:
        yield client


def account(client, name="student"):
    response = client.post("/auth/register", json={"username": name, "password": "a long test password"})
    assert response.status_code == 200, response.text
    return {"Authorization": "Bearer " + response.json()["token"]}


def upload(client, headers):
    response = client.post("/documents", headers=headers, files={"file": ("Biology.txt", MATERIAL.encode(), "text/plain")})
    assert response.status_code == 201, response.text
    return response.json()


def test_complete_loop_and_retry(client):
    headers = account(client)
    document = upload(client, headers)
    assert len(document["sections"]) == 2
    response = client.post(f'/documents/{document["id"]}/quizzes', headers=headers, json={"count": 10})
    assert response.status_code == 201, response.text
    quiz = response.json()
    quiz_id = quiz["id"]
    assert "answer" not in quiz["question"] and "quote" not in quiz["question"]
    assert client.get(f"/quizzes/{quiz_id}/results", headers=headers).status_code == 409
    assert client.post(f"/quizzes/{quiz_id}/next", headers=headers, json={"expected_index": 0}).status_code == 409
    for index in range(10):
        quiz = client.get(f"/quizzes/{quiz_id}", headers=headers).json()
        body = {"question_id": quiz["question"]["id"], "choice": 0}
        answered = client.post(f"/quizzes/{quiz_id}/answers", headers=headers, json=body)
        assert answered.status_code == 200, answered.text
        assert answered.json()["feedback"]["quote"]
        assert client.post(f"/quizzes/{quiz_id}/answers", headers=headers, json=body).json() == answered.json()
        assert client.post(f"/quizzes/{quiz_id}/answers", headers=headers, json={**body, "choice": 1}).status_code == 409
        next_body = {"expected_index": index}
        advanced = client.post(f"/quizzes/{quiz_id}/next", headers=headers, json=next_body).json()
        assert client.post(f"/quizzes/{quiz_id}/next", headers=headers, json=next_body).json() == advanced
    results = client.get(f"/quizzes/{quiz_id}/results", headers=headers).json()
    assert results["total"] == 10
    assert results["correct"] + results["incorrect"] == 10
    progress = client.get("/progress", headers=headers).json()
    assert progress["quizzes_completed"] == 1
    assert progress["questions_answered"] == 10
    retry = client.post(f"/quizzes/{quiz_id}/retry", headers=headers)
    assert retry.status_code == (201 if results["incorrect"] else 422)
    if results["incorrect"]:
        assert retry.json()["count"] == results["incorrect"]


def test_owner_boundaries_and_deletion(client):
    first = account(client, "first")
    second = account(client, "second")
    document = upload(client, first)
    doc_id = document["id"]
    quiz = client.post(f"/documents/{doc_id}/quizzes", headers=first, json={"count": 5}).json()
    assert client.get("/documents").status_code == 401
    assert client.get("/documents", headers=second).json() == []
    for path in [f"/documents/{doc_id}", f"/documents/{doc_id}/passages/p0", f'/quizzes/{quiz["id"]}', f'/quizzes/{quiz["id"]}/results']:
        assert client.get(path, headers=second).status_code == 404
    assert client.delete(f"/documents/{doc_id}", headers=second).status_code == 404
    assert client.delete(f"/documents/{doc_id}", headers=first).status_code == 204
    assert client.get(f'/quizzes/{quiz["id"]}', headers=first).status_code == 404
    upload(client, first)
    assert client.delete("/account", headers=first).status_code == 204
    assert client.get("/documents", headers=first).status_code == 401
    store = client.app.state.store
    assert store.db.execute("SELECT COUNT(*) FROM records").fetchone()[0] == 0


def test_resume_survives_server_restart(tmp_path):
    with TestClient(create_app(tmp_path)) as first:
        headers = account(first)
        doc = upload(first, headers)
        quiz = first.post(f'/documents/{doc["id"]}/quizzes', headers=headers, json={"count": 5}).json()
        first.post(f'/quizzes/{quiz["id"]}/answers', headers=headers, json={"question_id": quiz["question"]["id"], "choice": 0})
    with TestClient(create_app(tmp_path)) as second:
        resumed = second.get(f'/quizzes/{quiz["id"]}', headers=headers).json()
        assert resumed["index"] == 0
        assert resumed["feedback"] is not None
        assert len(second.get("/progress", headers=headers).json()["in_progress"]) == 1


def test_selection_boundaries():
    document = structure("Biology.txt", extract("Biology.txt", MATERIAL.encode()))
    section = document["sections"][0]["id"]
    selected = retrieve(document, {"section_ids": [section]})
    assert all(p["section_id"] == section for p in selected)
    assert all(p["section_id"] == section for p in selected)
    assert "Ecosystems" not in " ".join(p["text"] for p in selected)
    for bad in [{"section_ids": ["unknown"]}, {"page_start": 1}, {"page_start": 1, "page_end": 2}, {"topic": "quantum"}]:
        with pytest.raises(DocumentError):
            retrieve(document, bad)
    pdf = structure("book.pdf", ["Chapter 1 A\n" + MATERIAL.splitlines()[1], "Chapter 2 B\n" + MATERIAL.splitlines()[8]])
    assert {p["page"] for p in retrieve(pdf, {"page_start": 2, "page_end": 2})} == {2}


def test_chapter_includes_children_and_wrapped_lines():
    doc = structure("book.pdf", ["Chapter 1 Cells\nSection 1.1 Energy\nMitochondria produce usable\nenergy through cellular respiration.\nChapter 2 Ecology\nEcosystems include living organisms."])
    assert doc["passages"][0]["text"] == "Mitochondria produce usable energy through cellular respiration."
    selected = retrieve(doc, {"section_ids": [doc["sections"][0]["id"]]})
    assert len(selected) == 1 and "Mitochondria" in selected[0]["text"]
    assert len(retrieve(doc, {"topic": "quiz me on Cells"})) == 1
    with pytest.raises(DocumentError):
        retrieve(doc, {"topic": "what is the"})


def test_docx_heading_styles_are_preserved():
    document = Document()
    document.add_heading("Cell Biology", level=1)
    document.add_heading("Energy Production", level=2)
    document.add_paragraph("Mitochondria produce energy inside living cells.")
    buffer = io.BytesIO()
    document.save(buffer)
    result = structure("notes.docx", extract("notes.docx", buffer.getvalue()))
    assert result["sections"][1]["parent_id"] == result["sections"][0]["id"]
    assert result["sections"][0]["title"] == "Cell Biology"


@pytest.mark.parametrize("filename,data", [("file.exe", b"anything"), ("bad.txt", b"\xff"), ("empty.txt", b""), ("null.txt", b"a\x00b"), ("bad.pdf", b"hello"), ("bad.docx", b"PKinvalid")])
def test_invalid_documents(filename, data):
    with pytest.raises(DocumentError):
        extract(filename, data)


def test_docx_and_pdf_failure_modes():
    doc = Document()
    doc.add_paragraph("A real document with readable text.")
    buffer = io.BytesIO()
    doc.save(buffer)
    assert "readable text" in extract("valid.docx", buffer.getvalue())[0]
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentError, match="OCR"):
        extract("scan.pdf", buffer.getvalue())
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentError, match="Password"):
        extract("locked.pdf", buffer.getvalue())


def test_real_text_pdf_upload_preserves_pages(client):
    writer = PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    for sentence in ["Mitochondria produce usable energy inside living cells.", "Ecosystems contain organisms and their surrounding environment."]:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({sentence}) Tj ET".encode())
        page[NameObject("/Contents")] = stream
    buffer = io.BytesIO()
    writer.write(buffer)
    headers = account(client)
    response = client.post("/documents", headers=headers, files={"file": ("pages.pdf", buffer.getvalue(), "application/pdf")})
    assert response.status_code == 201, response.text
    doc = response.json()
    assert doc["page_count"] == 2 and doc["has_real_pages"]
    passage = client.get(f'/documents/{doc["id"]}/passages/p1', headers=headers).json()
    assert passage["page"] == 2 and "Ecosystems" in passage["text"]


@pytest.mark.parametrize("kind", ["multiple_choice", "true_false", "yes_no"])
def test_generation_contract_and_rejection(kind):
    passages = structure("notes.txt", [MATERIAL])["passages"]
    questions = generate_local(passages, 10, kind, "easy")
    assert len(questions) == 10
    questions[0]["quote"] = "invented source"
    with pytest.raises(GenerationError):
        validate(questions, passages, 10)
    with pytest.raises(GenerationError):
        generate_local(passages[:1], 10, kind, "easy")
    with pytest.raises(GenerationError):
        generate_local(passages, 10, kind, "hard")


def test_encryption_retention_and_session_revocation(client):
    headers = account(client)
    document = upload(client, headers)
    store = client.app.state.store
    payload = store.db.execute("SELECT payload FROM records").fetchone()[0]
    assert b"Mitochondria" not in payload
    store.db.execute("UPDATE records SET created=?", (time.time() - 31 * 86400,))
    store.db.commit()
    assert client.get("/documents", headers=headers).json() == []
    assert client.delete("/auth/session", headers=headers).status_code == 204
    assert client.get("/documents", headers=headers).status_code == 401


def test_missing_storage_key_never_silently_rekeys_database(tmp_path):
    store = Store(tmp_path)
    store.db.close()
    (tmp_path / "storage.key").unlink()
    with pytest.raises(RuntimeError, match="encryption key is missing"):
        Store(tmp_path)
    assert not (tmp_path / "storage.key").exists()


def test_upload_limits_and_validation(client, monkeypatch):
    headers = account(client)
    assert client.post("/documents", headers={**headers, "content-length": str(60 * 1024 * 1024)}, content=b"x").status_code == 413
    response = client.post("/documents", headers=headers, files={"file": ("bad.pdf", b"invalid")})
    assert response.status_code == 422
    doc = upload(client, headers)
    assert client.post(f'/documents/{doc["id"]}/quizzes', headers=headers, json={"count": 51}).status_code == 422


def test_authentication_and_rate_limit(client):
    headers = account(client)
    assert client.post("/auth/login", json={"username": "student", "password": "incorrect password"}).status_code == 401
    assert client.post("/auth/login", json={"username": "student", "password": "a long test password"}).status_code == 200
    for _ in range(8):
        response = client.post("/auth/login", json={"username": "student", "password": "incorrect password"})
    assert response.status_code == 429
