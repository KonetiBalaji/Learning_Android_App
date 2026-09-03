from io import BytesIO

from docx import Document
import pytest

from app.document_service import extract_text
from app.docx_text import extract_docx_body
from app.config import get_settings
from app.database import connection
from test_mastery_loop import registered_client


def encoded(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_paragraph_only_text_preserves_existing_separator_behavior():
    doc = Document()
    doc.add_paragraph('Before 😀')
    doc.add_paragraph('   ')
    doc.add_paragraph('After\nSecond line')
    assert extract_text(encoded(doc), '.docx') == 'Before 😀\nAfter\nSecond line'


def test_tables_are_in_body_order_and_cells_preserve_row_order():
    doc = Document()
    doc.add_paragraph('Before')
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip([cell for row in table.rows for cell in row.cells], ['Organ', 'Function', 'Leaf', 'Photosynthesis']):
        cell.text = value
    doc.add_paragraph('After')
    assert extract_text(encoded(doc), '.docx') == 'Before\nOrgan\tFunction\nLeaf\tPhotosynthesis\nAfter'


def test_nested_table_keeps_cell_paragraph_order():
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = 'Cell before'
    cell.add_table(rows=1, cols=1).cell(0, 0).text = 'Nested fact'
    cell.add_paragraph('Cell after')
    assert extract_text(encoded(doc), '.docx') == 'Cell before\nNested fact\nCell after'


@pytest.mark.parametrize('vertical', [False, True])
def test_merged_cells_do_not_duplicate_source(vertical):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(1, 0) if vertical else table.cell(0, 1))
    merged.text = 'Merged source'
    table.cell(1, 1).text = 'Other source'
    result = extract_text(encoded(doc), '.docx')
    assert result.count('Merged source') == 1
    assert result.count('Other source') == 1


def test_table_depth_limit_rejects_not_truncates(monkeypatch):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = 'Nested'
    monkeypatch.setattr('app.docx_text.MAX_TABLE_DEPTH', 1)
    with pytest.raises(ValueError, match='nesting'):
        extract_docx_body(doc)


def test_text_and_grid_limits_reject_not_truncate(monkeypatch):
    doc = Document()
    doc.add_paragraph('Long source content')
    monkeypatch.setattr('app.docx_text.MAX_TEXT_CHARACTERS', 5)
    with pytest.raises(ValueError, match='text'):
        extract_docx_body(doc)
    doc = Document()
    doc.add_table(rows=2, cols=2)
    monkeypatch.setattr('app.docx_text.MAX_GRID_CELLS', 3)
    with pytest.raises(ValueError, match='cell'):
        extract_docx_body(doc)


def test_table_only_upload_becomes_searchable_with_exact_stored_offsets(monkeypatch, tmp_path):
    client, headers = registered_client(monkeypatch, tmp_path)
    try:
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = 'Photosynthesis converts light energy into chemical energy within green plants.'
        table.cell(1, 0).text = 'Chlorophyll absorbs light, supporting the production of sugars from carbon dioxide and water.'
        result = client.post('/v1/documents', headers=headers, files={'file': ('tables.docx', encoded(doc))})
        assert result.status_code == 201, result.text
        uploaded = result.json()
        searched = client.post(f"/v1/documents/{uploaded['id']}/search", headers=headers,
            json={'chapter_ids': [c['id'] for c in uploaded['chapters']], 'query': 'chlorophyll'})
        assert searched.status_code == 200
        hit = searched.json()['matches'][0]
        with connection(get_settings().database_path) as db:
            stored = db.execute('SELECT extracted_text FROM documents WHERE id = ?', (uploaded['id'],)).fetchone()[0]
        assert hit['text'] == stored[hit['start_character']:hit['end_character']]
        assert stored[hit['match_start']:hit['match_end']] == 'Chlorophyll'
    finally:
        client.__exit__(None, None, None)
